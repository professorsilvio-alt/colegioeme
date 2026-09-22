import sys
import datetime
import json
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.shortcuts import get_object_or_404, redirect, render
from django.db import transaction
from django.views.decorators.http import require_POST

from ..models import AulaExtraProgramada, Disciplina, Professor, Turma
from ..utils import get_professor


def _extract_rows_from_excel(file_obj):
    import openpyxl
    workbook = openpyxl.load_workbook(file_obj, data_only=True)
    sheet = workbook.active
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        rows.append(row)
    return rows

def _extract_rows_from_word(file_obj):
    import docx
    doc = docx.Document(file_obj)
    rows = []
    for table in doc.tables:
        for i, row in enumerate(table.rows):
            if i == 0: continue # Skip header roughly
            row_data = [cell.text.strip() for cell in row.cells]
            rows.append(row_data)
    return rows

def _extract_rows_from_pdf(file_obj):
    import pdfplumber
    rows = []
    with pdfplumber.open(file_obj) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            for table in tables:
                for i, row in enumerate(table):
                    if i == 0: continue # Skip header
                    clean_row = [str(cell).replace('\n', ' ').strip() if cell else '' for cell in row]
                    rows.append(clean_row)
    return rows


@login_required
def upload_aulas_extras(request):
    prof = get_professor(request.user)
    if not prof or not prof.pode_editar_tudo:
        messages.error(request, 'Acesso restrito à administração e direção.')
        return redirect('dashboard')

    if request.method == 'POST':
        # PASSO 2: Confirmação Final
        if 'parsed_data_json' in request.POST:
            try:
                data_json = request.POST.get('parsed_data_json')
                valid_rows = json.loads(data_json)
                sucesso = 0
                
                with transaction.atomic():
                    for item in valid_rows:
                        # Re-fetch models securely
                        turma = Turma.objects.get(pk=item['turma_id'])
                        disc = Disciplina.objects.get(pk=item['disc_id'])
                        professor = Professor.objects.get(pk=item['prof_id'])
                        data_aula = datetime.datetime.strptime(item['data_str'], '%Y-%m-%d').date()
                        
                        obj, created = AulaExtraProgramada.objects.get_or_create(
                            data=data_aula,
                            turma=turma,
                            disciplina=disc,
                            professor=professor
                        )
                        if created:
                            sucesso += 1
                            
                messages.success(request, f"{sucesso} aulas extras registradas com sucesso!")
                return redirect('upload_aulas_extras')
            except Exception as e:
                messages.error(request, f"Erro ao confirmar os dados: {str(e)}")
                return redirect('upload_aulas_extras')

        # PASSO 1: Upload e Extração
        file_obj = request.FILES.get('arquivo_upload')
        usar_ia = request.POST.get('usar_ia') == '1'

        if not file_obj:
            messages.error(request, 'Por favor, selecione um arquivo.')
            return redirect('upload_aulas_extras')

        MAX_UPLOAD_SIZE = 10 * 1024 * 1024  # 10 MB
        if file_obj.size > MAX_UPLOAD_SIZE:
            messages.error(request, f'Arquivo muito grande (máximo {MAX_UPLOAD_SIZE // (1024*1024)} MB).')
            return redirect('upload_aulas_extras')

        filename = file_obj.name.lower()
        rows = []
        
        valid_rows = []
        erros = []
        
        if usar_ia:
            from core.services.ai_parser import extrair_aulas_extras_com_ia
            try:
                dados_ia = extrair_aulas_extras_com_ia(file_obj, filename, request.ano_letivo)
                for i, item in enumerate(dados_ia, start=1):
                    try:
                        # O item esperado tem: data, horario, turma, disciplina, professor
                        val_data = item.get('data', '')
                        val_turma = item.get('turma', '')
                        val_disc = item.get('disciplina', '')
                        val_prof = item.get('professor', '')
                        
                        if not all([val_data, val_turma, val_disc, val_prof]):
                            erros.append(f"Registro da IA {i}: Campos essenciais ausentes na extração.")
                            continue

                        # Tratamento da data
                        if '/' in val_data:
                            d, m, y = val_data.split('/')
                            if len(y) == 2: y = "20" + y
                            data_aula = datetime.date(int(y), int(m), int(d))
                        else:
                            data_aula = datetime.date.fromisoformat(val_data[:10])
                            
                        # Tratamento de Modelos
                        turma = Turma.objects.filter(codigo__iexact=str(val_turma).strip(), ano_letivo=request.ano_letivo, escola=request.escola).first()
                        if not turma:
                            erros.append(f"Registro IA {i}: Turma '{val_turma}' não encontrada.")
                            continue
                            
                        disciplina = Disciplina.objects.filter(nome__iexact=str(val_disc).strip()).first()
                        if not disciplina:
                            erros.append(f"Registro IA {i}: Disciplina '{val_disc}' não encontrada.")
                            continue
                            
                        professor = Professor.objects.filter(nome__iexact=str(val_prof).strip()).first()
                        if not professor:
                            professores = Professor.objects.filter(nome__icontains=str(val_prof).strip())
                            if professores.count() == 1:
                                professor = professores.first()
                            else:
                                erros.append(f"Registro IA {i}: Professor '{val_prof}' não encontrado unicamente.")
                                continue
                                
                        display_text = f"{data_aula.strftime('%d/%m/%Y')} | {item.get('horario','')} | Turma {turma.codigo} | {disciplina.nome} | {professor.nome}"
                        valid_rows.append({
                            'data_str': data_aula.isoformat(),
                            'turma_id': turma.id,
                            'disc_id': disciplina.id,
                            'prof_id': professor.id,
                            'display': display_text
                        })
                    except Exception as e:
                        erros.append(f"Registro IA {i}: Falha ao processar o dado - {str(e)}")
            except Exception as e:
                messages.error(request, f'Erro na leitura com IA: {str(e)}')
                return redirect('upload_aulas_extras')
                
        else:
            try:
                if filename.endswith('.xlsx'):
                    rows = _extract_rows_from_excel(file_obj)
                elif filename.endswith('.docx'):
                    rows = _extract_rows_from_word(file_obj)
                elif filename.endswith('.pdf'):
                    rows = _extract_rows_from_pdf(file_obj)
                else:
                    messages.error(request, 'Formato de arquivo inválido. Use .xlsx, .docx ou .pdf.')
                    return redirect('upload_aulas_extras')
            except Exception as e:
                messages.error(request, f'Erro ao ler o arquivo (Modo Padrão): Certifique-se de que é um formato válido com tabelas legíveis. Detalhe: {str(e)}')
                return redirect('upload_aulas_extras')

        if not usar_ia:
            # Processar linhas extraídas do modo padrão
            for i, row in enumerate(rows, start=2): # Start 2 just as indicative line number
                if not row or not any(row):
                    continue
                    
                if len(row) < 5:
                    erros.append(f"Linha {i}: Colunas insuficientes na tabela (esperado Data, Horário, Turma, Disc, Prof).")
                    continue
                    
                val_data = row[0]
                val_turma = row[2]
                val_disc = row[3]
                val_prof = row[4]
                
                if not all([val_data, val_turma, val_disc, val_prof]):
                    erros.append(f"Linha {i}: Campos essenciais em branco.")
                    continue

                # Parse Date
                data_aula = None
                if isinstance(val_data, datetime.datetime):
                    data_aula = val_data.date()
                else:
                    try:
                        val_data_str = str(val_data).strip()
                        if '/' in val_data_str:
                            d, m, y = val_data_str.split('/')
                            if len(y) == 2: y = "20" + y
                            data_aula = datetime.date(int(y), int(m), int(d))
                        else:
                            data_aula = datetime.date.fromisoformat(val_data_str[:10])
                    except Exception:
                        erros.append(f"Linha {i}: Data inválida '{val_data}'. Formato exigido: DD/MM/AAAA")
                        continue
                        
                # Match Turma
                turma = Turma.objects.filter(codigo__iexact=str(val_turma).strip(), ano_letivo=request.ano_letivo, escola=request.escola).first()
                if not turma:
                    erros.append(f"Linha {i}: Turma '{val_turma}' não encontrada no ano/escola atual.")
                    continue
                    
                # Match Disciplina
                disciplina = Disciplina.objects.filter(nome__iexact=str(val_disc).strip()).first()
                if not disciplina:
                    erros.append(f"Linha {i}: Disciplina '{val_disc}' não encontrada.")
                    continue
                    
                # Match Professor
                professor = Professor.objects.filter(nome__iexact=str(val_prof).strip()).first()
                if not professor:
                    professores = Professor.objects.filter(nome__icontains=str(val_prof).strip())
                    if professores.count() == 1:
                        professor = professores.first()
                    else:
                        erros.append(f"Linha {i}: Professor '{val_prof}' não encontrado de forma única.")
                        continue

                # Add to valid list
                display_text = f"{data_aula.strftime('%d/%m/%Y')} | {turma.codigo} | {disciplina.nome} | {professor.nome}"
                valid_rows.append({
                    'data_str': data_aula.isoformat(),
                    'turma_id': turma.id,
                    'disc_id': disciplina.id,
                    'prof_id': professor.id,
                    'display': display_text
                })

        if not valid_rows and not erros:
            messages.warning(request, "O sistema não encontrou nenhuma tabela válida no documento.")
            return redirect('upload_aulas_extras')

        context = {
            'preview': True,
            'valid_rows': valid_rows,
            'erros': erros,
            'parsed_json': json.dumps(valid_rows)
        }
        return render(request, 'core/upload_aulas_extras.html', context)

    return render(request, 'core/upload_aulas_extras.html')


@login_required
def gerenciar_aulas_extras(request):
    """Lista as aulas extras cadastradas com filtros por professor, turma e período,
    permitindo reatribuição individual ou em massa para outro professor."""
    from ..models import AnoLetivo
    prof = get_professor(request.user)
    if not prof or not prof.pode_editar_tudo:
        messages.error(request, 'Acesso restrito à administração e direção.')
        return redirect('dashboard')

    professores = Professor.objects.filter(escolas=request.escola).order_by('nome')
    turmas = Turma.objects.filter(
        ano_letivo=request.ano_letivo, escola=request.escola
    ).order_by('ordem_exibicao', 'codigo')

    filtro_prof = request.GET.get('professor', '')
    filtro_turma = request.GET.get('turma', '')
    filtro_data_inicio = request.GET.get('data_inicio', '')
    filtro_data_fim = request.GET.get('data_fim', '')

    qs = AulaExtraProgramada.objects.filter(
        turma__ano_letivo=request.ano_letivo,
        turma__escola=request.escola
    ).select_related('turma', 'disciplina', 'professor').order_by('-data', 'turma__codigo', 'professor__nome')

    if filtro_prof:
        qs = qs.filter(professor_id=filtro_prof)
    if filtro_turma:
        qs = qs.filter(turma_id=filtro_turma)
    if filtro_data_inicio:
        try:
            dt = datetime.datetime.strptime(filtro_data_inicio, '%Y-%m-%d').date()
            qs = qs.filter(data__gte=dt)
        except ValueError:
            pass
    if filtro_data_fim:
        try:
            dt = datetime.datetime.strptime(filtro_data_fim, '%Y-%m-%d').date()
            qs = qs.filter(data__lte=dt)
        except ValueError:
            pass

    filtro_aplicado = bool(filtro_prof or filtro_turma or filtro_data_inicio or filtro_data_fim)

    context = {
        'prof': prof,
        'aulas': qs if filtro_aplicado else AulaExtraProgramada.objects.none(),
        'professores': professores,
        'turmas': turmas,
        'disciplinas_todas': Disciplina.objects.all().order_by('nome'),
        'filtro_prof': filtro_prof,
        'filtro_turma': filtro_turma,
        'filtro_data_inicio': filtro_data_inicio,
        'filtro_data_fim': filtro_data_fim,
        'filtro_aplicado': filtro_aplicado,
        'total': qs.count() if filtro_aplicado else 0,
    }
    return render(request, 'core/gerenciar_aulas_extras.html', context)


@login_required
def aula_extra_editar(request, pk):
    """Edita os dados de uma única AulaExtraProgramada (troca professor, data, turma ou disciplina)."""
    from ..models import Disciplina as Disc
    prof = get_professor(request.user)
    if not prof or not prof.pode_editar_tudo:
        messages.error(request, 'Acesso restrito à administração e direção.')
        return redirect('dashboard')

    aula = get_object_or_404(AulaExtraProgramada, pk=pk)

    if request.method == 'POST':
        novo_prof_id = request.POST.get('professor')
        nova_turma_id = request.POST.get('turma')
        nova_disc_id = request.POST.get('disciplina')
        nova_data_str = request.POST.get('data')

        try:
            nova_data = datetime.datetime.strptime(nova_data_str, '%Y-%m-%d').date()
            novo_prof = Professor.objects.get(pk=novo_prof_id)
            nova_turma = Turma.objects.get(pk=nova_turma_id)
            nova_disc = Disc.objects.get(pk=nova_disc_id)

            # Verifica conflito unique_together antes de salvar
            conflito = AulaExtraProgramada.objects.filter(
                data=nova_data, turma=nova_turma, disciplina=nova_disc, professor=novo_prof
            ).exclude(pk=aula.pk).exists()

            if conflito:
                messages.error(request, 'Já existe uma aula extra com esses dados. Nenhuma alteração foi feita.')
            else:
                aula.data = nova_data
                aula.professor = novo_prof
                aula.turma = nova_turma
                aula.disciplina = nova_disc
                aula.save()
                messages.success(request, f'Aula extra atualizada com sucesso!')

        except Exception as e:
            messages.error(request, f'Erro ao salvar: {str(e)}')

        return redirect(request.POST.get('next', 'gerenciar_aulas_extras'))

    return redirect('gerenciar_aulas_extras')


@login_required
@require_POST
def aulas_extras_reatribuir(request):
    """Ação em massa: reatribui as aulas extras selecionadas de qualquer professor para um novo professor."""
    from ..models import Disciplina as Disc
    prof = get_professor(request.user)
    if not prof or not prof.pode_editar_tudo:
        messages.error(request, 'Acesso restrito à administração e direção.')
        return redirect('dashboard')

    ids = request.POST.getlist('aula_id')
    novo_prof_id = request.POST.get('novo_professor')
    next_url = request.POST.get('next', '')

    if not ids:
        messages.error(request, 'Nenhuma aula extra selecionada.')
        return redirect(_safe_redirect(next_url))

    if not novo_prof_id:
        messages.error(request, 'Selecione o professor de destino.')
        return redirect(_safe_redirect(next_url))

    try:
        novo_prof = Professor.objects.get(pk=novo_prof_id)
    except Professor.DoesNotExist:
        messages.error(request, 'Professor de destino não encontrado.')
        return redirect(_safe_redirect(next_url))

    sucesso = 0
    ignorados = 0

    with transaction.atomic():
        for aula in AulaExtraProgramada.objects.filter(pk__in=ids).select_related('turma', 'disciplina', 'professor'):
            if aula.professor_id == novo_prof.pk:
                # Já é o professor alvo — nada a fazer
                ignorados += 1
                continue

            conflito = AulaExtraProgramada.objects.filter(
                data=aula.data,
                turma=aula.turma,
                disciplina=aula.disciplina,
                professor=novo_prof
            ).exists()

            if conflito:
                ignorados += 1
            else:
                aula.professor = novo_prof
                aula.save()
                sucesso += 1

    if sucesso:
        msg = f'{sucesso} aula(s) extra(s) reatribuída(s) para {novo_prof.nome} com sucesso.'
        if ignorados:
            msg += f' {ignorados} não foram alteradas (já vinculadas ao professor ou conflito de duplicidade).'
        messages.success(request, msg)
    else:
        messages.warning(request, f'Nenhuma aula extra foi reatribuída. {ignorados} entradas já pertenciam ao professor ou apresentaram conflito de duplicidade.')

    return redirect(_safe_redirect(next_url))


def _safe_redirect(url):
    """Retorna a URL se for relativa e segura, caso contrário retorna a view padrão."""
    if url and url.startswith('/') and not url.startswith('//'):
        return url
    return 'gerenciar_aulas_extras'
